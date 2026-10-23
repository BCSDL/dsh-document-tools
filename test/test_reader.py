import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from document_reader import document
from worker import dispatch


class IndependentReaderTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.out = self.root / 'out'
        self.out.mkdir()

    def tearDown(self):
        self.temp.cleanup()

    def test_actual_pdf_pages_pixels_and_pagination_without_classroom(self):
        import fitz
        source = self.root / 'lesson.pdf'
        with fitz.open() as pdf:
            for text in ['Verified first page', 'Verified second page']:
                page = pdf.new_page()
                page.insert_text((50, 50), text)
            pdf.save(source)
        first = dispatch({'dataDirectory': str(self.out)}, {'path': str(source), 'limit': 1})
        second = dispatch({'dataDirectory': str(self.out)}, {'path': str(source), 'offset': first['nextOffset'], 'limit': 1})
        self.assertTrue(first['truncated'])
        self.assertFalse(second['truncated'])
        self.assertIn('first page', first['pages'][0]['text'])
        self.assertIn('second page', second['pages'][0]['text'])
        self.assertTrue(Path(first['pages'][0]['image']).is_file())

    def test_code_with_unfamiliar_extension_and_utf16(self):
        source = self.root / 'source.university_language'
        source.write_text('class Ball { field int x; }', encoding='utf-16')
        self.assertIn('class Ball', list(document(source, self.out))[0]['text'])

    def test_opaque_binary_is_not_claimed_readable(self):
        source = self.root / 'source.opaque'
        source.write_bytes(b'opaque\0binary')
        with self.assertRaisesRegex(ValueError, '二进制'):
            list(document(source, self.out))

    def test_docx_and_spreadsheet_are_independent(self):
        from docx import Document
        from openpyxl import Workbook
        source = self.root / 'notes.docx'
        doc = Document()
        doc.add_paragraph('Actual class notes')
        doc.save(source)
        self.assertIn('Actual class notes', list(document(source, self.out))[0]['text'])
        source = self.root / 'numbers.xlsx'
        book = Workbook()
        book.active.append(['formula', '=1+2'])
        book.save(source)
        book.close()
        self.assertIn('=1+2', list(document(source, self.out))[0]['text'])

    def test_one_shot_worker_exits_and_rejects_other_domain_operations(self):
        script = Path(__file__).resolve().parents[1] / 'backend/worker.py'
        result = subprocess.run([sys.executable, str(script)], input=json.dumps({'operation': 'start_file', 'args': {}}), text=True, capture_output=True, timeout=15, env={**os.environ, 'DSH_TOOL_CONFIG': json.dumps({'dataDirectory': str(self.out)})})
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('Unknown document operation', json.loads(result.stdout)['error'])


if __name__ == '__main__':
    unittest.main()
